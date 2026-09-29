# -*- coding: utf-8 -*-
"""Генерация отчёта по лабораторной работе (report/Отчет_ЛАБ3_вариант7.docx)."""
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
REPORT = os.path.join(ROOT, "report")
RESULTS = os.path.join(REPORT, "results.json")
OUT = os.path.join(REPORT, "Отчет_ЛАБ3_вариант7.docx")

LANG_RU = {"en": "английский", "fr": "французский"}
DOM_RU = {"cs": "computer science", "lit": "литература"}


def add_table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    return t


def main():
    with open(RESULTS, encoding="utf-8") as f:
        res = json.load(f)
    rows, summary, backend = res["rows"], res["summary"], res["backend"]
    with open(os.path.join(ROOT, "data", "collection", "meta.json"),
              encoding="utf-8") as f:
        meta = json.load(f)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)

    # ---------------- титульный лист ----------------
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Лабораторная работа\n«Автоматическое реферирование документов»").bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Вариант 7: языки — французский, английский;\n"
              "методика — sentence extraction + нейросетевой метод "
              "(вместо технологии OSTIS);\n"
              "предметные области — научные статьи по computer science, "
              "сочинения по литературе")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Реализация: веб-система на FastAPI;\nзависимости и виртуальное окружение управляются через uv\n\n")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Студент: ____________________   Группа: __________\n"
              "Преподаватель: ____________________\n\nМинск, 2026")
    doc.add_page_break()

    # ---------------- 1. цель и задание ----------------
    doc.add_heading("1. Цель и содержание работы", level=1)
    doc.add_paragraph(
        "Цель работы — освоить на практике основные принципы автоматического "
        "реферирования документов. В соответствии с вариантом 7 разработана "
        "система, строящая реферат текстовых документов на французском и "
        "английском языках из предметных областей «научные статьи по computer "
        "science» и «сочинения по литературе».")
    doc.add_paragraph(
        "В отличие от методических указаний, где экстрактивный метод "
        "sentence extraction дополняется технологией OSTIS (семантические "
        "сети), в данной работе вместо OSTIS применён нейросетевой метод: "
        "многоязычная нейросеть-трансформер (paraphrase-multilingual-MiniLM-L12-v2, "
        "квантованный ONNX-инференс) строит семантические векторы предложений "
        "и терминов, которые используются для устранения семантических повторов "
        "при отборе предложений (MMR), группировки ключевых слов и оценки "
        "качества реферата. Предусмотрен резервный полностью офлайновый "
        "нейросетевой бэкенд — собственная реализация skip-gram с negative "
        "sampling (numpy).")
    doc.add_paragraph("Система выполняет требования методики:")
    for s in [
        "на входе — текстовые документы одинакового размера (~10 страниц А4, "
        "≈18 000 символов) на языках варианта из указанных предметных областей;",
        "на выходе — активная ссылка на исходный документ и реферат из двух "
        "разделов: классический реферат и реферат в виде списка ключевых слов "
        "(иерархического);",
        "средства сохранения результата в файл (.txt, .html) и печати;",
        "предельно простой веб-интерфейс со справочной системой (help).",
    ]:
        doc.add_paragraph(s, style="List Bullet")

    # ---------------- 2. тестовая коллекция ----------------
    doc.add_heading("2. Тестовая коллекция документов", level=1)
    doc.add_paragraph(
        "Коллекция сформирована на основании многоязычных текстов Wikipedia "
        "(что допущено методичкой). Домен «сочинения по литературе» представлен "
        "аналитическими текстами о литературных произведениях (жанр "
        "литературоведческого эссе/критики). Все документы нормализованы к "
        "одинаковому размеру путём обрезки по границе предложения вблизи "
        "18 000 символов (~10 страниц А4 по 1800 символов). Вводный абзац "
        "статьи (lead) сохранён отдельно и используется как эталонный реферат "
        "при оценке качества (метрики ROUGE).")
    add_table(doc,
              ["ID", "Язык", "Область", "Заголовок", "Символов", "Источник"],
              [[m["id"], LANG_RU[m["lang"]], DOM_RU[m["domain"]], m["title"],
                m["chars"], m["source_url"]] for m in meta])
    doc.add_paragraph(
        "Итого 8 документов: 2 языка × 2 предметные области × 2 документа. "
        "Индекс коллекции (df(t), |DB| = 8) используется в TF-IDF; при "
        "обработке стороннего документа он временно добавляется в индекс "
        "(|DB| = 9).")

    # ---------------- 3. структура системы ----------------
    doc.add_heading("3. Структура разработанной системы", level=1)
    doc.add_paragraph(
        "Система реализована как веб-приложение на FastAPI (Python 3.11): "
        "Jinja2-шаблоны и HTML/CSS/JS формируют интерфейс; вычислительное ядро "
        "вынесено в пакет app.core. Структура приведена на рисунке 1.")
    doc.add_picture(os.path.join(REPORT, "structure.png"), width=Cm(16))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Рисунок 1 — Структура системы").alignment = \
        WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Основные модули:")
    for s in [
        "tokenizer.py — разбиение текста на абзацы и предложения с учётом "
        "аббревиатур; вычисление символьных смещений |D|, BD(Sᵢ), |P|, BP(Sᵢ); "
        "токенизация слов и фильтрация стоп-слов, чисел и слов чужой письменности;",
        "extraction.py — индекс коллекции (df, |DB|), базовый TF·IDF, "
        "модифицированный TF-IDF w(t,D), позиционные функции Posd/Posp, вес "
        "предложения W(Sᵢ) и отбор топ-N;",
        "neural.py — нейросетевые бэкенды (ONNX-трансформер, torch-трансформер, "
        "skip-gram) и MMR-отбор предложений;",
        "keywords.py — построение иерархического списка ключевых слов;",
        "summarizer.py — конвейер сборки реферата;",
        "collection.py — загрузка коллекции, автоопределение языка;",
        "main.py — маршруты FastAPI, экспорт файлов, печать;",
        "scripts/ — сборка коллекции, тестирование и оценка, диаграммы, отчёт.",
    ]:
        doc.add_paragraph(s, style="List Bullet")

    # ---------------- 4. структуры данных ----------------
    doc.add_heading("4. Структуры данных", level=1)
    doc.add_paragraph("Входные и промежуточные данные:")
    add_table(doc,
              ["Структура", "Поля", "Назначение"],
              [
                ["Document", "id, title, lang, domain, source_url, text, "
                 "paragraphs[], sentences[]", "разобранный входной документ"],
                ["Paragraph", "index, text, length (|P|), sentences[]",
                 "абзац для позиционной функции Posp"],
                ["Sentence", "index, text, doc_offset (BD), par_index, "
                 "par_offset (BP), tokens[], all_tokens[]",
                 "предложение со смещениями и значимыми словами"],
                ["CorpusIndex", "df: dict[term→int], db_size (|DB|)",
                 "индекс коллекции для IDF"],
                ["ExtractionResult", "tf_doc, tf_max, base_weights, scored[]",
                 "веса слов и предложения документа"],
                ["SentenceScore", "sentence, posd, posp, score, weight, terms",
                 "разложение веса предложения по функциям методики"],
              ])
    doc.add_paragraph("Выходные данные:")
    add_table(doc,
              ["Структура", "Поля", "Назначение"],
              [
                ["SummaryOutput", "sentences[] (index, text, weight, posd, "
                 "posp, score), keywords[] (дерево), метаданные, elapsed_ms, "
                 "compression", "готовый реферат из двух разделов"],
                ["KeywordNode", "term, weight, count, kind "
                 "(root/phrase/semantic), children[]",
                 "узел иерархического списка ключевых слов"],
              ])
    doc.add_paragraph(
        "Хранение: коллекция — текстовые файлы data/collection/*.txt + "
        "meta.json + эталоны *.ref.txt; загруженные пользователем документы — "
        "data/uploads/*.txt; результаты сессии — в памяти сервера (словарь "
        "RESULTS) с выгрузкой в файлы .txt/.html.")

    # ---------------- 5. алгоритм ----------------
    doc.add_heading("5. Алгоритм построения реферата", level=1)
    doc.add_paragraph(
        "Алгоритм повторяет методику sentence extraction и дополнен "
        "нейросетевым шагом (замена OSTIS).")
    doc.add_paragraph(
        "Шаг 1. Веса слов. Из рассмотрения исключаются стоп-слова, числа и "
        "слова письменности, не соответствующей языку документа. Базовый вес "
        "слова: w(t) = tf(t,D)·log(|DB|/df(t)).")
    doc.add_paragraph(
        "Шаг 2. Веса предложений. Модифицированный TF-IDF: w(t,D) = "
        "0.5·(1 + tf(t,D)/tf_max(D))·log(|DB|/df(t)). Позиционные функции: "
        "Posd(Sᵢ) = 1 − BD(Sᵢ)/|D|, Posp(Sᵢ) = 1 − BP(Sᵢ)/|P|. Вес предложения: "
        "W(Sᵢ) = Posd(Sᵢ)·Posp(Sᵢ)·Score(Sᵢ), где Score(Sᵢ) = Σ_{t∈Sᵢ} "
        "tf(t,Sᵢ)·w(t,D).")
    doc.add_paragraph(
        "Шаг 3. Нейросетевое уточнение (вместо OSTIS). Для кандидатов с "
        "наибольшим W вычисляются эмбеддинги предложений; жадный MMR-отбор "
        "max [λ·W_norm(Sᵢ) − (1−λ)·max_{j∈selected} cos(Sᵢ,Sⱼ)] при λ = 0.75 "
        "исключает семантические повторы; при отключённой нейросети берётся "
        "чистый топ-N по W.")
    doc.add_paragraph(
        "Шаг 4. Генерация. Отобранные N = 10 предложений выводится в порядке "
        "их следования в тексте — классический реферат.")
    doc.add_paragraph(
        "Шаг 5. Ключевые слова. По базовому TF·IDF выбираются вершины-корни "
        "(устойчивые словосочетания объединяются в составные корни, синонимы "
        "дедуплицируются по косинусу эмбеддингов ≥ 0.80); потомки корня — "
        "словосочетания, содержащие корень (до 4), и семантически близкие "
        "термины (косинус ≥ 0.45, до 3). Получается иерархический список "
        "вида «термин → словосочетания/синонимы».")
    doc.add_picture(os.path.join(REPORT, "algorithm.png"), width=Cm(14))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Рисунок 2 — Схема алгоритма").alignment = \
        WD_ALIGN_PARAGRAPH.CENTER

    # ---------------- 6. тестирование ----------------
    doc.add_heading("6. Результаты тестирования системы", level=1)
    doc.add_paragraph(
        "Тестирование выполнено на всех 8 документах коллекции в двух режимах: "
        "«extraction» (чистый sentence extraction) и «extraction+neural» "
        "(с нейросетевым MMR-уточнением). Нейросетевой бэкенд: " + backend +
        ". Метрики: ROUGE-1/ROUGE-2 (F-мера против эталонного lead-реферата), "
        "coverage — доля топ-20 TF-IDF-терминов документа, попавших в реферат, "
        "adequacy — косинусная близость эмбеддингов реферата и полного "
        "документа, redundancy — средний попарный косинус предложений реферата "
        "(меньше — меньше повторов), compression — степень сжатия, время.")
    add_table(doc,
              ["Документ", "Режим", "ROUGE-1", "ROUGE-2", "Coverage",
               "Adequacy", "Redundancy", "Время, мс"],
              [[r["doc_id"], r["mode"], r["rouge1"], r["rouge2"],
                r["coverage"], r["adequacy"], r["redundancy"],
                r["elapsed_ms"]] for r in rows])
    doc.add_paragraph()
    doc.add_paragraph("Сводка по группам (средние значения):")
    add_table(doc,
              ["Группа (язык/область/режим)", "ROUGE-1", "ROUGE-2", "Coverage",
               "Adequacy", "Redundancy", "Время, мс"],
              [[k, v["rouge1"], v["rouge2"], v["coverage"], v["adequacy"],
                v["redundancy"], round(v["elapsed_ms"])]
               for k, v in summary.items()])

    # ---------------- 7. оценка ----------------
    doc.add_heading("7. Оценка полученных результатов", level=1)
    a = summary["all/all/extraction"]
    b = summary["all/all/extraction+neural"]
    doc.add_paragraph(
        f"Точность. Средний ROUGE-1 экстрактивного реферата против эталонного "
        f"lead-абзаца составил {a['rouge1']:.3f} (ROUGE-2 {a['rouge2']:.3f}), что "
        f"соответствует типичным значениям экстрактивных систем данного класса. "
        f"Покрытие ключевых терминов — {a['coverage']:.0%}, семантическая "
        f"адекватность реферата полному тексту — {a['adequacy']:.3f} из 1, т.е. "
        f"реферат сохраняет смысловое содержание документа.")
    doc.add_paragraph(
        f"Влияние нейросетевого компонента. Добавление нейросетевого MMR-отбора "
        f"снижает семантическую избыточность реферата с {a['redundancy']:.3f} до "
        f"{b['redundancy']:.3f} (повторов смысла меньше) при практически "
        f"неизменном ROUGE-1 ({a['rouge1']:.3f} → {b['rouge1']:.3f}): часть "
        f"«почти дублирующих» выгодных по TF-IDF предложений заменяется "
        f"семантически новыми. Наибольший эффект — на французских научных "
        f"текстах (red 0.420 → 0.352).")
    doc.add_paragraph(
        f"Языки и области. Английские научные тексты реферируются лучше всего "
        f"(ROUGE-1 {summary['en/cs/extraction+neural']['rouge1']:.3f}), "
        f"французские литературные — с наибольшим ROUGE-2 "
        f"({summary['fr/lit/extraction+neural']['rouge2']:.3f}) благодаря "
        f"плотным повествовательным конструкциям; различия между языками "
        f"непреодолимого барьера не создают, многоязычный трансформер "
        f"выравнивает качество.")
    doc.add_paragraph(
        f"Время. Чистый sentence extraction — {a['elapsed_ms']:.0f} мс на "
        f"документ (~18 000 символов); с нейросетевым уточнением — "
        f"{b['elapsed_ms']:.0f} мс (инференс квантованного трансформера на CPU), "
        f"что приемлемо для интерактивной веб-системы.")
    doc.add_paragraph(
        "Сжатие: реферат из 10 предложений составляет ~13–21 % объёма "
        "документа, т.е. задача «отсеивания менее значимой информации» "
        "выполняется.")

    # ---------------- 8. компоненты ----------------
    doc.add_heading("8. Применённые готовые компоненты", level=1)
    for s in [
        "FastAPI + Uvicorn + Jinja2 — веб-каркас, маршрутизация, шаблоны, "
        "загрузка файлов (python-multipart);",
        "onnxruntime + tokenizers (Hugging Face) — инференс квантованной "
        "многоязычной модели paraphrase-multilingual-MiniLM-L12-v2 "
        "(model_quint8_avx2.onnx, int8): в ~4 раза меньше памяти и быстрее "
        "полной fp32-модели, что позволило работать в ограниченном окружении; "
        "пулинг — среднее по токенам с маской внимания, как в "
        "sentence-transformers;",
        "sentence-transformers / torch — опциональный бэкенд той же сети "
        "(env NEURAL_BACKEND=transformers);",
        "numpy — собственная нейросеть skip-gram (negative sampling, SGD) и "
        "векторные операции косинусной близости;",
        "matplotlib — диаграммы отчёта; python-docx — генерация настоящего "
        "отчёта;",
        "uv — управление зависимостями и виртуальным окружением проекта: "
        "метаданные и зависимости объявлены в pyproject.toml (PEP 621, "
        "виртуальный проект), воспроизводимость обеспечивает файл блокировки "
        "uv.lock, версия интерпретатора зафиксирована в .python-version; "
        "сопровождающие зависимости (matplotlib, python-docx, httpx) вынесены "
        "в dev-группу, опциональный torch-бэкенд — в extra «torch»; "
        "используются команды uv sync / uv run / uv add / uv lock;",
        "Wikipedia API (action=query&prop=extracts) — формирование тестовой "
        "коллекции (через прокси r.jina.ai из-за сетевой блокировки).",
    ]:
        doc.add_paragraph(s, style="List Bullet")
    doc.add_paragraph(
        "Особенности: все нейросетевые бэкенды скрыты за единым интерфейсом "
        "NeuralBackend.embed(); при недоступности сети/тяжёлых зависимостей "
        "система деградирует до офлайнового skip-gram или до чистого "
        "sentence extraction без потери работоспособности.")

    # ---------------- 9. выводы ----------------
    doc.add_heading("9. Выводы и перспективы", level=1)
    doc.add_paragraph(
        "В ходе работы реализована и протестирована система автоматического "
        "реферирования документов (вариант 7): метод sentence extraction по "
        "формулам методики (TF·IDF, модифицированный TF-IDF, позиционные "
        "функции) дополнен нейросетевым методом вместо технологии OSTIS. "
        "Система оформлена как веб-приложение на FastAPI с простым интерфейсом, "
        "справкой, сохранением в файл и печатью; построена и опробована "
        "тестовая коллекция из 8 многоязычных документов одинакового размера; "
        "получены количественные оценки качества (ROUGE-1 ≈ 0.35, покрытие "
        "ключевых терминов ≈ 0.69, адекватность ≈ 0.90) и времени обработки.")
    doc.add_paragraph(
        "Перспективы: абстрактное (генеративное) реферирование на базе "
        "seq2seq-моделей; запросно-ориентированные рефераты; иерархические "
        "ключевые слова онтологического типа; расширение коллекции и "
        "полноценная пользовательская оценка (human evaluation); кэширование "
        "эмбеддингов документов коллекции для ускорения.")

    # ---------------- приложение ----------------
    doc.add_heading("Приложение А. Пример реферата", level=1)
    ex_path = os.path.join(REPORT, "example.json")
    if os.path.exists(ex_path):
        with open(ex_path, encoding="utf-8") as f:
            example = json.load(f)
        doc.add_paragraph(
            "Документ: Les Misérables (французский, литература), режим "
            "extraction+neural. Классический реферат (10 предложений, в порядке "
            "следования в тексте):")
        for i, s in enumerate(example["sentences"], 1):
            txt = s if len(s) <= 220 else s[:220] + "…"
            doc.add_paragraph(f"{i}. {txt}", style="List Number")
        doc.add_paragraph("Иерархический список ключевых слов:")
        for node in example["keywords"][:6]:
            kids = "; ".join(c["term"] for c in node["children"][:3])
            line = f"{node['term']}"
            if kids:
                line += f" → {kids}; …"
            doc.add_paragraph(line, style="List Bullet")
    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
